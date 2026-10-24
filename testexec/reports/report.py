"""Test raporu üretimi: HTML (Jinja2), PDF (Qt WebEngine) ve Excel (openpyxl)."""
from __future__ import annotations

import base64
import io
from pathlib import Path

import numpy as np
from jinja2 import Environment, FileSystemLoader, select_autoescape

from ..db.models import TestRun

_env = Environment(loader=FileSystemLoader(Path(__file__).parent / "templates"),
                   autoescape=select_autoescape(["html"]))


def _chart_png_b64(step) -> str | None:
    """Adımın ham ölçüm serisini limit bantlarıyla birlikte PNG olarak döner."""
    if not step.series:
        return None
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    y = np.frombuffer(step.series, dtype="<f8")
    fig, ax = plt.subplots(figsize=(6.4, 2.2), dpi=130)
    ax.plot(y, color="#2563eb", lw=1.6)
    if step.low is not None:
        ax.axhline(step.low, color="#dc2626", ls="--", lw=1)
    if step.high is not None:
        ax.axhline(step.high, color="#dc2626", ls="--", lw=1)
    ax.set_title(step.name, fontsize=9, loc="left")
    ax.set_xlabel("Örnek", fontsize=8)
    ax.set_ylabel(step.unit, fontsize=8)
    ax.tick_params(labelsize=8)
    ax.grid(alpha=.25)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return base64.b64encode(buf.getvalue()).decode()


def render_html(run: TestRun, include_traffic: bool = False) -> str:
    steps = [dict(s=s, chart=_chart_png_b64(s)) for s in sorted(run.steps, key=lambda s: s.idx)]
    traffic = []
    if include_traffic:
        for t in sorted(run.traffic, key=lambda t: t.ts):
            traffic.append((datetime_str(t.ts), t.channel, t.direction, t.data.hex(" ")))
    return _env.get_template("report.html").render(run=run, steps=steps, traffic=traffic)


def datetime_str(ts: float) -> str:
    from datetime import datetime
    return datetime.fromtimestamp(ts).strftime("%H:%M:%S.%f")[:-3]


def save_html(run: TestRun, path: str | Path, include_traffic: bool = False) -> Path:
    p = Path(path)
    p.write_text(render_html(run, include_traffic), encoding="utf-8")
    return p


def save_pdf(run: TestRun, path: str | Path, include_traffic: bool = False) -> Path:
    """HTML'i Qt WebEngine ile PDF'e basar (QApplication gerekir)."""
    from PySide6.QtCore import QEventLoop, QTimer, QUrl
    from PySide6.QtWebEngineCore import QWebEnginePage
    p = Path(path)
    html = render_html(run, include_traffic)
    page = QWebEnginePage()
    loop = QEventLoop()
    ok = {"v": False}

    def on_pdf(_path, success):
        ok["v"] = success
        loop.quit()

    page.pdfPrintingFinished.connect(on_pdf)
    page.loadFinished.connect(lambda s: page.printToPdf(str(p)) if s else loop.quit())
    page.setHtml(html, QUrl("about:blank"))
    QTimer.singleShot(60000, loop.quit)
    loop.exec()
    if not ok["v"]:
        raise RuntimeError("PDF oluşturulamadı")
    return p


def save_excel(runs: list[TestRun], path: str | Path) -> Path:
    """Seçili koşuları iki sayfalık Excel'e yazar: Koşular ve Adımlar."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Koşular"
    ws.append(["ID", "Sekans", "Seri No", "Operatör", "İstasyon", "Başlangıç", "Bitiş", "Sonuç"])
    ws2 = wb.create_sheet("Adımlar")
    ws2.append(["Koşu ID", "Seri No", "Adım", "Sonuç", "Değer", "Birim", "Min", "Max", "Süre (s)", "Mesaj"])
    for r in runs:
        ws.append([r.id, r.sequence, r.serial, r.operator, r.station, r.started, r.finished, r.status])
        for s in sorted(r.steps, key=lambda s: s.idx):
            ws2.append([r.id, r.serial, s.name, s.status,
                        s.value_num if s.value_num is not None else s.value_text,
                        s.unit, s.low, s.high, round(s.duration, 3), s.message])
    colors = {"PASS": "C6F6D5", "FAIL": "FED7D7", "ERROR": "FEEBC8"}
    for sheet, col in ((ws, 8), (ws2, 4)):
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F2937")
        for row in sheet.iter_rows(min_row=2):
            color = colors.get(str(row[col - 1].value))
            if color:
                row[col - 1].fill = PatternFill("solid", fgColor=color)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 18
        sheet.freeze_panes = "A2"
    p = Path(path)
    wb.save(p)
    return p