"""CSV/Excel'den DUT mesaj kataloğu içe aktarma, şablon dışa aktarma ve otomatik sekans üretimi."""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Union

import yaml

from .catalog import Catalog, CatalogError, FieldDef, MessageDef

TEMPLATE_COLUMNS = [
    "Message", "ID", "Direction", "Response", "Group", "Description",
    "Field", "Offset", "Size", "Type", "Default", "Scale", "Unit", "Enum", "Bits"
]


def _parse_kv_pairs(text: str) -> dict[str, int]:
    """'NORMAL=0;TEST=1' veya 'OK:0, ERR:1' metnini sözlüğe çevirir."""
    out: dict[str, int] = {}
    if not text or not str(text).strip():
        return out
    for part in str(text).replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        sep = "=" if "=" in part else ":"
        if sep in part:
            k, v = part.split(sep, 1)
            out[k.strip()] = int(v.strip(), 0)
    return out


def _read_rows(path: Path) -> list[dict[str, str]]:
    ext = path.suffix.lower()
    if ext == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return [{k.strip(): (v.strip() if v else "") for k, v in row.items() if k}
                    for row in csv.DictReader(f)]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return []
        headers = [str(c).strip() if c is not None else "" for c in header_row]
        out: list[dict[str, str]] = []
        for r in rows_iter:
            if not r or all(c is None or str(c).strip() == "" for c in r):
                continue
            d = {}
            for h, val in zip(headers, r):
                if h:
                    d[h] = "" if val is None else str(val).strip()
            out.append(d)
        return out
    raise CatalogError(f"Desteklenmeyen dosya uzantısı: {ext} (.csv veya .xlsx kullanın)")


def import_catalog(path: Union[str, Path], dut_name: str = "ImportedDUT",
                   byte_order: str = "little") -> Catalog:
    """CSV veya Excel tablosundan Catalog nesnesi oluşturur."""
    p = Path(path)
    rows = _read_rows(p)
    cat = Catalog(name=dut_name or p.stem, byte_order=byte_order)
    grouped: dict[str, dict[str, Any]] = {}

    for row in rows:
        mname = row.get("Message") or row.get("Mesaj") or ""
        if not mname:
            continue
        if mname not in grouped:
            mid_str = row.get("ID", "0")
            mid = int(mid_str, 0) if isinstance(mid_str, str) else int(mid_str)
            grouped[mname] = {
                "name": mname,
                "id": mid,
                "direction": (row.get("Direction") or "tx").lower(),
                "response": row.get("Response") or None,
                "group": row.get("Group") or "",
                "description": row.get("Description") or "",
                "fields": [],
            }
        fname = row.get("Field") or row.get("Alan") or ""
        if fname:
            f_type = (row.get("Type") or "uint8").lower()
            offset = int(row.get("Offset") or 0, 0)
            size_raw = row.get("Size") or ""
            size: Union[int, str, None] = None
            if size_raw.lower() == "rest":
                size = "rest"
            elif size_raw:
                size = int(size_raw, 0)
            scale = float(row.get("Scale") or 1.0)
            unit = row.get("Unit") or ""
            default_raw = row.get("Default") or ""
            default: Any = None
            if default_raw != "":
                if f_type in ("bytes", "ascii"):
                    default = default_raw
                elif f_type.startswith("float"):
                    default = float(default_raw)
                else:
                    try:
                        default = int(default_raw, 0)
                    except ValueError:
                        default = default_raw
            grouped[mname]["fields"].append(FieldDef(
                name=fname, offset=offset, size=size, type=f_type,
                enum=_parse_kv_pairs(row.get("Enum", "")),
                bits=_parse_kv_pairs(row.get("Bits", "")),
                scale=scale, unit=unit, default=default,
            ))

    for mdata in grouped.values():
        cat.add(MessageDef(**mdata))
    cat.validate()
    return cat


def save_catalog_yaml(cat: Catalog, path: Union[str, Path]) -> Path:
    """Catalog nesnesini YAML profil dosyası olarak kaydeder."""
    d: dict[str, Any] = {"dut": cat.name, "byte_order": cat.byte_order, "messages": {}}
    for m in cat.messages.values():
        md: dict[str, Any] = {"id": f"0x{m.id:02X}", "direction": m.direction}
        if m.response:
            md["response"] = m.response
        if m.group:
            md["group"] = m.group
        if m.description:
            md["description"] = m.description
        if m.length is not None:
            md["length"] = m.length
        if m.fields:
            flist = []
            for f in m.fields:
                fd: dict[str, Any] = {"name": f.name, "offset": f.offset, "type": f.type}
                if f.is_rest or f.type in ("bytes", "ascii", "bits"):
                    fd["size"] = f.size
                if f.default is not None:
                    fd["default"] = f.default.hex(" ") if isinstance(f.default, bytes) else f.default
                if f.scale != 1.0:
                    fd["scale"] = f.scale
                if f.bias != 0.0:
                    fd["bias"] = f.bias
                if f.unit:
                    fd["unit"] = f.unit
                if f.enum:
                    fd["enum"] = f.enum
                if f.bits:
                    fd["bits"] = f.bits
                if f.description:
                    fd["description"] = f.description
                flist.append(fd)
            md["fields"] = flist
        d["messages"][m.name] = md
    if cat.presets:
        d["presets"] = cat.presets
    p = Path(path)
    p.write_text(yaml.safe_dump(d, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return p


def export_template_excel(path: Union[str, Path], cat: Catalog | None = None) -> Path:
    """Örnek (veya mevcut katalogdan) Excel mesaj tanım şablonu üretir."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Messages"
    ws.append(TEMPLATE_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")

    if cat:
        for m in cat.messages.values():
            if not m.fields:
                ws.append([m.name, f"0x{m.id:02X}", m.direction, m.response or "", m.group, m.description,
                           "", "", "", "", "", "", "", "", ""])
            else:
                for idx, f in enumerate(m.fields):
                    enum_s = ";".join(f"{k}={v}" for k, v in f.enum.items())
                    bits_s = ";".join(f"{k}={v}" for k, v in f.bits.items())
                    def_s = "" if f.default is None else (f.default.hex(" ") if isinstance(f.default, bytes) else str(f.default))
                    ws.append([
                        m.name, f"0x{m.id:02X}" if idx == 0 else "",
                        m.direction if idx == 0 else "",
                        (m.response or "") if idx == 0 else "",
                        m.group if idx == 0 else "",
                        m.description if idx == 0 else "",
                        f.name, f.offset, f.size, f.type, def_s,
                        f.scale if f.scale != 1.0 else "", f.unit, enum_s, bits_s
                    ])
    else:
        ws.append(["INIT_REQ", "0x10", "tx", "INIT_RSP", "init", "DUT başlatma",
                   "mode", 0, 1, "uint8", "NORMAL", "", "", "NORMAL=0;TEST=1", ""])
        ws.append(["INIT_REQ", "", "", "", "", "",
                   "init_params", 1, 3, "bytes", "00 00 00", "", "", "", ""])
        ws.append(["INIT_REQ", "", "", "", "", "",
                   "timeout_ms", 4, 2, "uint16", "500", "", "ms", "", ""])
        ws.append(["INIT_RSP", "0x90", "rx", "", "init", "Başlatma yanıtı",
                   "status", 0, 1, "uint8", "", "", "", "OK=0;ERR=1", ""])
        ws.append(["INIT_RSP", "", "", "", "", "",
                   "temp", 1, 2, "int16", "", "0.1", "°C", "", ""])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16
    ws.freeze_panes = "A2"
    p = Path(path)
    wb.save(p)
    return p


def generate_sequence_from_catalog(cat: Catalog, transport: str = "dut",
                                   group: str = "", preset: str = "",
                                   dut_profile_path: str = "") -> dict[str, Any]:
    """Katalogdaki tüm TX mesajları için otomatik test sekansı sözlüğü üretir."""
    steps = []
    for m in cat.search(group=group):
        if m.direction not in ("tx", "both"):
            continue
        step: dict[str, Any] = {
            "name": f"{m.name} Testi",
            "type": "dut_message",
            "transport": transport,
            "message": m.name,
        }
        if preset and preset in cat.presets and m.name in cat.presets[preset]:
            step["preset"] = preset
        if m.response and m.response in cat.messages:
            rm = cat.get(m.response)
            expect: dict[str, Any] = {}
            for rf in rm.fields:
                if rf.enum and "OK" in rf.enum:
                    expect[rf.name] = "OK"
            if expect:
                step["expect"] = expect
        steps.append(step)
    seq: dict[str, Any] = {
        "sequence": f"{cat.name} Otomatik Mesaj Testi" + (f" ({group})" if group else ""),
        "steps": steps,
    }
    if dut_profile_path:
        seq["dut_profile"] = dut_profile_path
    return seq